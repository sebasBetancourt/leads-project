"""Importa a Postgres una exportación de la pestaña Leads (.xlsx) para no repetir discover y enrich.

Sirve para llevar los leads de una máquina a otra o para recuperar una base perdida. Lo que no está
en la pestaña (dirección, fila original de gosom, de dónde salió cada contacto) queda vacío; el
place_id se recupera del enlace de Google Maps, así un `discover` posterior no duplica negocios.

Los leads entran con el estado que tenían (sin pasar por la máquina de estados: no es una
transición, es una restauración) y con un evento `imported`. Uno que ya existe no se toca.
"""

import re
from dataclasses import dataclass
from pathlib import Path

from leadbot.core.db import session_scope
from leadbot.enrichment.phone import normalize_co
from leadbot.leads.models import Company, ContactChannel, Event
from leadbot.leads.repository import LeadRepository
from leadbot.leads.status import Channel, LeadStatus
from leadbot.sheets.sync import LEADS_TAB

PLACE_ID_RE = re.compile(r"!19s(ChIJ[\w-]+)")
CHANNEL_COLUMNS = {
    "email": Channel.EMAIL,
    "whatsapp": Channel.WHATSAPP,
    "instagram": Channel.INSTAGRAM,
    "facebook": Channel.FACEBOOK,
    "linkedin": Channel.LINKEDIN,
}


@dataclass
class ImportStats:
    rows: int = 0
    created: int = 0
    existing: int = 0
    skipped: int = 0  # filas sin nombre, sector o ciudad
    contacts: int = 0


def place_id_from_link(link: str) -> str | None:
    match = PLACE_ID_RE.search(link or "")
    return match.group(1) if match else None


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(value, kind):
    text = _text(value).replace(",", ".")
    try:
        return kind(float(text)) if text else None
    except ValueError:
        return None


def read_rows(path: Path, tab: str = LEADS_TAB) -> list[dict[str, str]]:
    """Filas de la pestaña como dict columna -> texto."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[tab] if tab in workbook.sheetnames else workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [_text(h).lower() for h in next(rows, [])]
    return [dict(zip(header, row, strict=False)) for row in rows if any(v not in (None, "") for v in row)]


def import_rows(rows: list[dict]) -> ImportStats:
    stats = ImportStats(rows=len(rows))
    with session_scope() as session:
        repo = LeadRepository(session)
        for row in rows:
            name, sector, city = _text(row.get("empresa")), _text(row.get("sector")), _text(row.get("ciudad"))
            if not (name and sector and city):
                stats.skipped += 1
                continue

            link = _text(row.get("google maps"))
            place_id = place_id_from_link(link)
            phone = normalize_co(_text(row.get("teléfono"))) if _text(row.get("teléfono")) else None
            if repo.find_existing(place_id, phone) is not None:
                stats.existing += 1
                continue

            status = _text(row.get("estado")).upper()
            company = Company(
                place_id=place_id,
                name=name,
                sector=sector,
                city=city,
                phone_e164=phone,
                website=_text(row.get("web")) or None,
                google_category=_text(row.get("categoría")) or None,
                rating=_number(row.get("rating"), float),
                reviews_count=_number(row.get("reseñas"), int),
                source_query="importado de Sheets",
                raw={"link": link} if link else {},
                status=LeadStatus(status) if status in LeadStatus.__members__ else LeadStatus.NEW,
                score=_number(row.get("score"), int),
            )
            session.add(company)
            session.flush()

            for column, channel in CHANNEL_COLUMNS.items():
                for value in dict.fromkeys(v.strip() for v in _text(row.get(column)).splitlines() if v.strip()):
                    session.add(ContactChannel(company_id=company.id, channel=channel, value=value, source="import"))
                    stats.contacts += 1
            session.add(
                Event(
                    company_id=company.id,
                    kind="imported",
                    data={"from": "xlsx", "sheet_id": _text(row.get("id")), "status": company.status.value},
                )
            )
            stats.created += 1
    return stats


def import_xlsx(path: Path) -> ImportStats:
    return import_rows(read_rows(path))
